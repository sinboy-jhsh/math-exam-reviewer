"""
考卷雙模態輸入智慧解析器 (ExamParser)
支援解析：
1. Word 格式 (.docx)：精準提取段落、題號凸排、<m:oMath> 原生方程式與出題範圍
2. PDF 格式 (.pdf)：採用 pypdf 提取頁面文字並重構題目結構
3. 純文字/Markdown (.txt, .md)
"""

import os
import re
from dataclasses import dataclass, field
from typing import List, Optional, Dict, Any
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

@dataclass
class ExamQuestion:
    index: int
    q_type: str                   # 'choice', 'fill', 'comp'
    score: float
    stem: str
    options: List[str] = field(default_factory=list)
    claimed_answer: Optional[str] = None
    scope: Optional[str] = None
    has_figure: bool = False
    raw_omml_xml: Optional[str] = None

@dataclass
class ExamData:
    file_path: str
    file_type: str                # 'docx', 'pdf', 'text'
    school: str = "新北市立錦和高級中學"
    grade: str = "G9"             # G7, G8, G9, G10, G11, G12
    grade_text: str = "國中部九年級"
    title: str = "段考試題"
    scope_desc: str = "段考命題範圍"
    questions: List[ExamQuestion] = field(default_factory=list)
    total_score: float = 0.0

class ExamParser:
    @staticmethod
    def parse(file_path: str) -> ExamData:
        """主入口：依副檔名自動分流解析"""
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".docx":
            return ExamParser.parse_docx(file_path)
        elif ext == ".pdf":
            return ExamParser.parse_pdf(file_path)
        else:
            return ExamParser.parse_text(file_path)

    @staticmethod
    def _extract_options(text: str) -> List[str]:
        """
        強健提取 (A), (B), (C), (D) 選項，避免數值坐標如 (3, 9) 誤觸截斷
        """
        parts = re.split(r"(?=\([A-Da-d]\)|(?<!\w)[A-Da-d][\.\、]\s*)", text)
        options = []
        for part in parts:
            p = part.strip()
            if re.match(r"^\([A-Da-d]\)|^([A-Da-d][\.\、])", p):
                options.append(p)
        return options

    @staticmethod
    def _extract_paragraph_text(p) -> str:
        r"""完整提取段落文字，包含 Word 原生 EQ 欄位代碼 (如 \x\to(AB) 轉 AB, \F(a, b) 轉 a/b) 與 OMML 數學式"""
        parts = []
        in_instr = False
        instr_buffer = []
        for elem in p._p.iter():
            tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
            if tag == "t":
                if not in_instr:
                    parts.append(elem.text or "")
            elif tag == "instrText":
                instr_buffer.append(elem.text or "")
            elif tag == "fldChar":
                fldType = elem.attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", "")
                if fldType == "begin":
                    in_instr = True
                    instr_buffer = []
                elif fldType in ["separate", "end"]:
                    in_instr = False
                    if instr_buffer:
                        eq_str = "".join(instr_buffer).strip()
                        # 解析線段 \x\to(AB) -> AB
                        m_x = re.search(r"\\x\\to\((.*?)\)", eq_str, re.IGNORECASE)
                        # 解析分數 \F(a, b) -> a/b
                        m_f = re.search(r"\\F\((.*?),(.*?)\)", eq_str, re.IGNORECASE)
                        if m_x:
                            parts.append(m_x.group(1))
                        elif m_f:
                            parts.append(f"{m_f.group(1).strip()}/{m_f.group(2).strip()}")
                        else:
                            parts.append(eq_str)
                        instr_buffer = []
        return "".join(parts).strip()

    @staticmethod
    def parse_docx(file_path: str) -> ExamData:
        doc = docx.Document(file_path)
        exam = ExamData(file_path=file_path, file_type="docx")

        # 1. 識別卷頭資訊
        for p in doc.paragraphs[:8]:
            txt = ExamParser._extract_paragraph_text(p)
            if not txt:
                continue
            if "錦和" in txt:
                exam.school = "新北市立錦和高級中學"
            if "七年級" in txt or "國七" in txt:
                exam.grade = "G7"; exam.grade_text = "國中部七年級"
            elif "八年級" in txt or "國八" in txt:
                exam.grade = "G8"; exam.grade_text = "國中部八年級"
            elif "九年級" in txt or "國九" in txt:
                exam.grade = "G9"; exam.grade_text = "國中部九年級"
            elif "高一" in txt or "十年級" in txt:
                exam.grade = "G10"; exam.grade_text = "高中部一年級"
            elif "高二" in txt or "十一年級" in txt:
                exam.grade = "G11"; exam.grade_text = "高中部二年級"
            elif "高三" in txt or "十二年級" in txt:
                exam.grade = "G12"; exam.grade_text = "高中部三年級"

            m_range = re.search(r"﹝命題範圍：(.*?)﹞|範圍[：:](.*)", txt)
            if m_range:
                exam.scope_desc = (m_range.group(1) or m_range.group(2) or "").strip()

            if "段考" in txt or "試題" in txt:
                exam.title = txt

        # 2. 將全卷展開為連續行串流，精確處理跨行與同段落多行題目
        lines = []
        for p in doc.paragraphs:
            p_text = ExamParser._extract_paragraph_text(p)
            for sub in p_text.splitlines():
                s = sub.strip()
                if s:
                    lines.append(s)

        current_type = "choice"
        current_score = 4.0
        questions = []
        active_q: Optional[ExamQuestion] = None
        choice_counter = 0
        fill_counter = 0
        comp_counter = 0

        for line in lines:
            # 大題切換判斷
            if any(k in line for k in ["第一部分", "選擇題", "單選題", "單一選擇題"]):
                current_type = "choice"
                m_sc = re.search(r"每題\s*(\d+)\s*分", line)
                if m_sc: current_score = float(m_sc.group(1))
                continue
            elif any(k in line for k in ["第二部分", "填充題"]):
                current_type = "fill"
                m_sc = re.search(r"每[題格]\s*(\d+)\s*分", line)
                if m_sc: current_score = float(m_sc.group(1))
                continue
            elif any(k in line for k in ["第三部分", "非選擇題", "計算證明題", "計算題", "非選題"]):
                current_type = "comp"
                m_sc = re.search(r"每[題小題]\s*(\d+)\s*分", line)
                if m_sc: current_score = float(m_sc.group(1))
                continue

            is_new_q = False
            q_idx = 0
            content = ""

            if current_type == "choice":
                # 支援 1.( )、( )、1.、以及帶圖示前綴的題號
                m_bnum = re.search(r"(\d+)[\.\、\s]*[\(（]\s*[\)）]\s*(.*)", line)
                m_brk = re.match(r"^[\(（]\s*[\)）]\s*(.*)", line)
                m_num = re.match(r"^(\d+)[\.\、\s]\s*(.*)", line)
                if m_bnum:
                    is_new_q = True
                    q_idx = int(m_bnum.group(1))
                    choice_counter = q_idx
                    content = m_bnum.group(2).strip()
                elif m_brk:
                    is_new_q = True
                    choice_counter += 1
                    q_idx = choice_counter
                    content = m_brk.group(1).strip()
                elif m_num and not line.startswith("(A)") and not line.startswith("A."):
                    is_new_q = True
                    q_idx = int(m_num.group(1))
                    choice_counter = q_idx
                    content = m_num.group(2).strip()
            elif current_type == "fill":
                if "(2)" in line and "(1)" not in line:
                    is_new_q = True
                    fill_counter += 1
                    q_idx = fill_counter
                    content = line
                else:
                    m_num = re.match(r"^(\d+)[\.\、\s]\s*(.*)", line)
                    starters = ["有一", "已知", "若", "如圖", "設", "在", "一", "有", "求", "將", "某", "△", "知△"]
                    is_start = (m_num is not None) or any(line.startswith(s) for s in starters)
                    if m_num:
                        is_new_q = True
                        q_idx = int(m_num.group(1))
                        fill_counter = q_idx
                        content = m_num.group(2).strip()
                    elif is_start and not line.startswith("(A)") and not line.startswith("(B)"):
                        is_new_q = True
                        fill_counter += 1
                        q_idx = fill_counter
                        content = line
            elif current_type == "comp":
                m_num = re.match(r"^(\d+)[\.\、\s]\s*(.*)", line)
                if m_num:
                    is_new_q = True
                    q_idx = int(m_num.group(1))
                    comp_counter = q_idx
                    content = m_num.group(2).strip()

            if is_new_q:
                if active_q:
                    questions.append(active_q)

                # 提取出題範圍 (支援全半形括號及減號，如 (1-1) 或 (1－2))
                scope_str = None
                m_sc_tag = re.search(r"[\(（]([1-9][\-\－][1-9][a-zA-Z]?)[\)）]", content)
                if m_sc_tag:
                    scope_str = m_sc_tag.group(1)
                    content = content[:m_sc_tag.start()] + content[m_sc_tag.end():]

                # 提取選項
                options = ExamParser._extract_options(content)
                if options:
                    first_opt_idx = content.find(options[0])
                    stem_clean = content[:first_opt_idx].strip()
                else:
                    stem_clean = content.strip()

                active_q = ExamQuestion(
                    index=q_idx,
                    q_type=current_type,
                    score=current_score,
                    stem=stem_clean,
                    options=options,
                    scope=scope_str
                )
            else:
                if active_q:
                    # 檢查是否包含出題範圍標籤
                    m_sc_tag = re.search(r"[\(（]([1-9][\-\－][1-9][a-zA-Z]?)[\)）]", line)
                    if m_sc_tag and not active_q.scope:
                        active_q.scope = m_sc_tag.group(1)
                        line = line[:m_sc_tag.start()] + line[m_sc_tag.end():]

                    # 檢查是否為續接選項
                    opts_found = ExamParser._extract_options(line)
                    if opts_found:
                        active_q.options.extend(opts_found)
                    else:
                        active_q.stem += " " + line.strip()

        if active_q:
            questions.append(active_q)

        # 嘗試從同目錄的解答檔自動載入答案
        ExamParser._pair_answer_key(questions, file_path)

        exam.questions = questions
        exam.total_score = sum(q.score for q in questions)
        return exam

    @staticmethod
    def _extract_cell_text(cell) -> str:
        r"""提取儲存格文字，完整解析 Word EQ 欄位代碼中的分數 \F(a, b)、根號 \r(n) 與線段 \x\to(AB)"""
        parts = []
        in_instr = False
        instr_buf = []
        for elem in cell._tc.iter():
            tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
            if tag == "t":
                if not in_instr:
                    parts.append(elem.text or "")
            elif tag == "instrText":
                instr_buf.append(elem.text or "")
            elif tag == "fldChar":
                fType = elem.attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", "")
                if fType == "begin":
                    in_instr = True
                    instr_buf = []
                elif fType in ["separate", "end"]:
                    in_instr = False
                    if instr_buf:
                        eq = "".join(instr_buf).strip()
                        m_f = re.search(r"\\F\(\s*(.*?)\s*,\s*(.*?)\s*\)", eq, re.IGNORECASE)
                        if m_f:
                            num = m_f.group(1).replace(r"\r(", "√(").replace(r"\R(", "√(")
                            num = re.sub(r"√\(\s*(\d+)\s*\)", r"√\1", num).strip()
                            den = m_f.group(2).strip()
                            parts.append(f"{num}/{den}")
                        else:
                            m_x = re.search(r"\\x\\to\((.*?)\)", eq, re.IGNORECASE)
                            if m_x:
                                parts.append(m_x.group(1))
                            else:
                                parts.append(eq)
                        instr_buf = []
        return "".join(parts).strip()

    @staticmethod
    def _pair_answer_key(questions: List[ExamQuestion], exam_path: str):
        """若同目錄存在解答卷 (.docx)，自動抽取標準答案配對"""
        dir_name = os.path.dirname(os.path.abspath(exam_path))
        ans_file = None
        for f in os.listdir(dir_name):
            if any(k in f for k in ["解答", "答案"]) and f.endswith(".docx") and not f.startswith("~$"):
                ans_file = os.path.join(dir_name, f)
                break

        if not ans_file or not os.path.exists(ans_file):
            return

        try:
            doc_ans = docx.Document(ans_file)
            ans_map = {}
            if len(doc_ans.tables) >= 1:
                t_c = doc_ans.tables[0]
                if len(t_c.rows) >= 2:
                    for col_idx, cell in enumerate(t_c.rows[0].cells):
                        q_num = ExamParser._extract_cell_text(cell)
                        ans = ExamParser._extract_cell_text(t_c.rows[1].cells[col_idx])
                        if q_num.isdigit():
                            ans_map[f"choice_{q_num}"] = ans

            if len(doc_ans.tables) >= 2:
                t_f = doc_ans.tables[1]
                if len(t_f.rows) >= 4 and len(t_f.rows[0].cells) == 5:
                    for j in range(5):
                        v1 = ExamParser._extract_cell_text(t_f.rows[1].cells[j])
                        v2 = ExamParser._extract_cell_text(t_f.rows[3].cells[j])
                        ans_map[f"fill_{j+1}"] = v1
                        ans_map[f"fill_{j+6}"] = v2

            if len(doc_ans.tables) >= 3:
                t_comp = doc_ans.tables[2]
                for j, cell in enumerate(t_comp.rows[0].cells):
                    ans_map[f"comp_{j+1}"] = ExamParser._extract_cell_text(cell)

            for q in questions:
                key = f"{q.q_type}_{q.index}"
                if key in ans_map:
                    q.claimed_answer = ans_map[key]
        except Exception:
            pass

    @staticmethod
    def parse_pdf(file_path: str) -> ExamData:
        """使用 pypdf 解析 PDF 考卷內容"""
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        exam = ExamData(file_path=file_path, file_type="pdf")

        full_text = ""
        for page in reader.pages:
            t = page.extract_text()
            if t: full_text += t + "\n"

        lines = [line.strip() for line in full_text.splitlines() if line.strip()]

        current_type = "choice"
        current_score = 4.0
        questions = []
        active_q: Optional[ExamQuestion] = None

        for line in lines:
            if "錦和" in line: exam.school = "新北市立錦和高級中學"
            if "七年級" in line: exam.grade = "G7"; exam.grade_text = "國中部七年級"
            elif "八年級" in line: exam.grade = "G8"; exam.grade_text = "國中部八年級"
            elif "九年級" in line: exam.grade = "G9"; exam.grade_text = "國中部九年級"

            if any(k in line for k in ["選擇題", "單一選擇題"]):
                current_type = "choice"
                continue
            elif any(k in line for k in ["填充題"]):
                current_type = "fill"
                continue
            elif any(k in line for k in ["非選擇題", "計算證明題"]):
                current_type = "comp"
                continue

            m_num = re.match(r"^(\d+)[\.\、\s]\s*(.*)", line)
            if m_num:
                if active_q: questions.append(active_q)
                q_idx = int(m_num.group(1))
                content = m_num.group(2).strip()

                opts = ExamParser._extract_options(content)
                if len(opts) >= 3:
                    first_idx = content.find(opts[0])
                    stem = content[:first_idx].strip()
                else:
                    stem = content

                active_q = ExamQuestion(
                    index=q_idx,
                    q_type=current_type,
                    score=current_score,
                    stem=stem,
                    options=opts
                )
            else:
                if active_q and ("(A)" in line or "(B)" in line or "A." in line):
                    opts_found = ExamParser._extract_options(line)
                    if opts_found: active_q.options.extend(opts_found)
                elif active_q:
                    active_q.stem += " " + line

        if active_q: questions.append(active_q)

        exam.questions = questions
        exam.total_score = sum(q.score for q in questions)
        return exam

    @staticmethod
    def parse_text(file_path: str) -> ExamData:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

        exam = ExamData(file_path=file_path, file_type="text")
        lines = [line.strip() for line in content.splitlines() if line.strip()]

        current_type = "choice"
        current_score = 4.0
        questions = []
        active_q = None

        for line in lines:
            if "七年級" in line: exam.grade = "G7"
            elif "八年級" in line: exam.grade = "G8"
            elif "九年級" in line: exam.grade = "G9"

            if "選擇題" in line: current_type = "choice"; continue
            elif "填充題" in line: current_type = "fill"; continue
            elif "非選" in line or "計算" in line: current_type = "comp"; continue

            m = re.match(r"^(\d+)[\.\、]\s*(.*)", line)
            if m:
                if active_q: questions.append(active_q)
                active_q = ExamQuestion(
                    index=int(m.group(1)),
                    q_type=current_type,
                    score=current_score,
                    stem=m.group(2).strip()
                )
            else:
                if active_q and "(A)" in line:
                    opts = re.findall(r"\([A-D]\)[^\(]*", line)
                    active_q.options.extend([o.strip() for o in opts])
                elif active_q:
                    active_q.stem += " " + line

        if active_q: questions.append(active_q)
        exam.questions = questions
        exam.total_score = sum(q.score for q in questions)
        return exam
