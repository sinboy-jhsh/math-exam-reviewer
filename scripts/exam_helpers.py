import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

MATH_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

def set_page_margins(doc, top_cm=1.0, bottom_cm=1.0, left_cm=1.0, right_cm=1.0):
    """依照錦和高中規範設定頁邊界（預設全為 1.0 cm）"""
    for sec in doc.sections:
        sec.top_margin = Inches(top_cm / 2.54)
        sec.bottom_margin = Inches(bottom_cm / 2.54)
        sec.left_margin = Inches(left_cm / 2.54)
        sec.right_margin = Inches(right_cm / 2.54)

def setup_footer_page_number(doc):
    """在頁尾產生居中的〔第 X 頁,共 Y 頁〕動態頁碼"""
    for sec in doc.sections:
        footer = sec.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ""
        r1 = p.add_run("〔第")
        set_run_font(r1, font_name="標楷體", font_size=10)
        fld1 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p._p.append(fld1)
        r2 = p.add_run("頁,共")
        set_run_font(r2, font_name="標楷體", font_size=10)
        fld2 = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
        p._p.append(fld2)
        r3 = p.add_run("頁〕")
        set_run_font(r3, font_name="標楷體", font_size=10)

def set_run_font(run, font_name="標楷體", ascii_font="Times New Roman", font_size=11, bold=False, italic=False, underline=False, color=None):
    """設定中英文與字級標準：中文字型標楷體、英數字 Times New Roman"""
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    run.font.size = Pt(font_size)
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{ascii_font}" w:hAnsi="{ascii_font}" w:eastAsia="{font_name}"/>')
    rPr.append(rFonts)

# === Word 原生 OMML 數學方程式生成引擎 ===
def m_run(text, italic=False, font_size=10):
    """建立 OMML 數學文字節點 <m:r>，變數採 Cambria Math 斜體、數字與運算子正體"""
    sty = '<m:rPr><m:sty m:val="p"/></m:rPr>' if not italic else ''
    esc = str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    sz_val = int(font_size * 2)
    return f'<m:r>{sty}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math" w:eastAsia="標楷體"/><w:sz w:val="{sz_val}"/></w:rPr><m:t>{esc}</m:t></m:r>'

def m_frac(num_xml, den_xml):
    """建立原生分數 <m:f>：num_xml / den_xml"""
    return f'<m:f><m:num>{num_xml}</m:num><m:den>{den_xml}</m:den></m:f>'

def m_sqrt(base_xml, deg_xml=None):
    """建立原生根號 <m:rad>"""
    if deg_xml is None:
        return f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{base_xml}</m:e></m:rad>'
    return f'<m:rad><m:radPr><m:degHide m:val="0"/></m:radPr><m:deg>{deg_xml}</m:deg><m:e>{base_xml}</m:e></m:rad>'

def m_sup(base_xml, sup_xml):
    """建立原生上標/次方 <m:sSup>：base^{sup}"""
    return f'<m:sSup><m:e>{base_xml}</m:e><m:sup>{sup_xml}</m:sup></m:sSup>'

def m_sub(base_xml, sub_xml):
    """建立原生下標 <m:sSub>：base_{sub}"""
    return f'<m:sSub><m:e>{base_xml}</m:e><m:sub>{sub_xml}</m:sub></m:sSub>'

def m_bar(base_xml, pos="top"):
    r"""建立幾何線段頂線 <m:bar>：\overline{AB}"""
    return f'<m:bar><m:barPr><m:pos m:val="{pos}"/></m:barPr><m:e>{base_xml}</m:e></m:bar>'

def m_cases(*eq_xml_list):
    """建立聯立方程組左大括號 <m:d> + <m:eqArr>"""
    rows = "".join([f'<m:e>{eq}</m:e>' for eq in eq_xml_list])
    return f'<m:d><m:dPr><m:begChr m:val="{{"/><m:endChr m:val=""/></m:dPr><m:e><m:eqArr>{rows}</m:eqArr></m:e></m:d>'

def append_omath(paragraph, inner_omml_xml):
    """將組裝好的 OMML 字串轉為原生 <m:oMath> 物件附加至 Word 段落中"""
    omath_str = f'<m:oMath {MATH_NS}>{inner_omml_xml}</m:oMath>'
    paragraph._p.append(parse_xml(omath_str))

def split_top_level(s: str, op_chars: list) -> list:
    """在括號最外層，依據運算子群分割字串，保留運算子"""
    res = []
    depth = 0
    cur = []
    i = 0
    while i < len(s):
        c = s[i]
        if c == '(':
            depth += 1
            cur.append(c)
            i += 1
        elif c == ')':
            depth -= 1
            cur.append(c)
            i += 1
        elif depth == 0:
            matched_op = None
            for op in op_chars:
                if s[i:i+len(op)] == op:
                    matched_op = op
                    break
            if matched_op:
                res.append((''.join(cur), matched_op))
                cur = []
                i += len(matched_op)
            else:
                cur.append(c)
                i += 1
        else:
            cur.append(c)
            i += 1
    res.append((''.join(cur), None))
    return res

def is_matched(s: str) -> bool:
    depth = 0
    for c in s:
        if c == '(':
            depth += 1
        elif c == ')':
            depth -= 1
            if depth < 0:
                return False
    return depth == 0

def parse_math_to_omml(expr: str, font_size=10) -> str:
    """
    將數學算式字串轉換為 Word 原生 OMML XML。
    嚴格遵循運算子優先順序：關係運算符 -> 連比號 -> 加減 -> 乘法 -> 除法/分式 -> 根式/次方。
    """
    expr = expr.strip()
    if not expr:
        return ""

    # 前置符號標準化替換
    expr = expr.replace("=>", "⇒")
    expr = expr.replace("＝", "=")
    expr = expr.replace("~", "∼")
    expr = expr.replace("//", "∥")
    expr = expr.replace(" * ", " × ")
    expr = expr.replace("*", " × ")

    # 1. 關係/推導符號 (=, ⇒, ∼, ∥)
    rel_splits = split_top_level(expr, ["=", "⇒", "∼", "∥"])
    if len(rel_splits) > 1:
        parts = []
        for term, op in rel_splits:
            if term.strip():
                parts.append(parse_math_to_omml(term.strip(), font_size=font_size))
            if op:
                parts.append(m_run(f" {op} ", italic=False, font_size=font_size))
        return "".join(parts)

    # 2. 連比號 ':'
    ratio_splits = split_top_level(expr, [":"])
    if len(ratio_splits) > 1:
        parts = []
        for term, op in ratio_splits:
            if term.strip():
                parts.append(parse_add_sub(term.strip(), font_size=font_size))
            if op:
                parts.append(m_run(" : ", italic=False, font_size=font_size))
        return "".join(parts)

    return parse_add_sub(expr, font_size=font_size)

def parse_add_sub(expr: str, font_size=10) -> str:
    """處理加減運算 (+, -)"""
    expr = expr.strip()
    if not expr:
        return ""

    if expr.startswith("(") and expr.endswith(")") and is_matched(expr[1:-1]):
        if not re.search(r'\)\^[0-9a-zA-Z]+$', expr):
            return m_run("(", italic=False, font_size=font_size) + parse_math_to_omml(expr[1:-1], font_size) + m_run(")", italic=False, font_size=font_size)

    add_sub_splits = []
    depth = 0
    cur = []
    for idx, c in enumerate(expr):
        if c == '(':
            depth += 1
            cur.append(c)
        elif c == ')':
            depth -= 1
            cur.append(c)
        elif depth == 0 and c in ['+', '-'] and idx > 0 and expr[idx-1] not in ['+', '-', '*', '×', '/', '(', '=']:
            add_sub_splits.append((''.join(cur), c))
            cur = []
        else:
            cur.append(c)
    add_sub_splits.append((''.join(cur), None))

    if len(add_sub_splits) > 1:
        parts = []
        for term, op in add_sub_splits:
            if term.strip():
                parts.append(parse_mul(term.strip(), font_size=font_size))
            if op:
                parts.append(m_run(f" {op} ", italic=False, font_size=font_size))
        return "".join(parts)

    return parse_mul(expr, font_size=font_size)

def parse_mul(expr: str, font_size=10) -> str:
    """處理乘法 (×, *)"""
    expr = expr.strip()
    if not expr:
        return ""

    mul_splits = split_top_level(expr, ["×", "*"])
    if len(mul_splits) > 1:
        parts = []
        for term, op in mul_splits:
            if term.strip():
                parts.append(parse_div(term.strip(), font_size=font_size))
            if op:
                parts.append(m_run(" × ", italic=False, font_size=font_size))
        return "".join(parts)

    return parse_div(expr, font_size=font_size)

def parse_div(expr: str, font_size=10) -> str:
    """處理分式 (/)"""
    expr = expr.strip()
    if not expr:
        return ""

    div_splits = split_top_level(expr, ["/"])
    if len(div_splits) == 2:
        numer = div_splits[0][0].strip()
        denom = div_splits[1][0].strip()
        if numer.startswith("(") and numer.endswith(")") and is_matched(numer[1:-1]):
            numer = numer[1:-1].strip()
        if denom.startswith("(") and denom.endswith(")") and is_matched(denom[1:-1]):
            denom = denom[1:-1].strip()
        return m_frac(parse_math_to_omml(numer, font_size), parse_math_to_omml(denom, font_size))

    return parse_atom(expr, font_size=font_size)

def parse_atom(term: str, font_size=10) -> str:
    """處理基本元素：根號、次方、單項式、字母變數"""
    term = term.strip()
    if not term:
        return ""

    # 1. 根號：√(expr)
    sqrt_match = re.search(r'√\(([^)]+)\)', term)
    if sqrt_match:
        start, end = sqrt_match.span()
        before = term[:start]
        radicand = sqrt_match.group(1)
        after = term[end:]
        return parse_atom(before, font_size) + m_sqrt(parse_math_to_omml(radicand, font_size)) + parse_atom(after, font_size)

    # 根號：√[0-9a-zA-Z]+
    sqrt_simple = re.search(r'√([0-9a-zA-Z]+)', term)
    if sqrt_simple:
        start, end = sqrt_simple.span()
        before = term[:start]
        radicand = sqrt_simple.group(1)
        after = term[end:]
        return parse_atom(before, font_size) + m_sqrt(parse_atom(radicand, font_size)) + parse_atom(after, font_size)

    # 2. 次方：base^exp
    sup_match = re.search(r'(\([^\)]+\)|[0-9a-zA-Z]+)\^([0-9a-zA-Z]+|\([^\)]+\))', term)
    if sup_match:
        start, end = sup_match.span()
        before = term[:start]
        base = sup_match.group(1)
        exp = sup_match.group(2)
        after = term[end:]
        if base.startswith("(") and base.endswith(")") and is_matched(base[1:-1]):
            base_xml = m_run("(", italic=False, font_size=font_size) + parse_math_to_omml(base[1:-1], font_size) + m_run(")", italic=False, font_size=font_size)
        else:
            base_xml = parse_atom(base, font_size)
        if exp.startswith("(") and exp.endswith(")") and is_matched(exp[1:-1]):
            exp = exp[1:-1]
        exp_xml = parse_atom(exp, font_size)
        return parse_atom(before, font_size) + m_sup(base_xml, exp_xml) + parse_atom(after, font_size)

    # 3. 變數與常數拆分
    tokens = re.split(r'([A-Za-z]+|[0-9]+|[\+\-\×\(\),°△∠]+|\s+)', term)
    res = []
    for t in tokens:
        if not t:
            continue
        if re.match(r'^[A-Za-z]+$', t):
            res.append(m_run(t, italic=True, font_size=font_size))
        else:
            res.append(m_run(t, italic=False, font_size=font_size))
    return "".join(res)

def append_mixed_text(paragraph, text: str, font_size=9.5, bold=False, default_color=None):
    """將一段混合中文與數學運算式的字串，精準渲染為包含 OMML 原生方程式的 Word 段落"""
    if not text:
        return

    text = text.replace("＝", "=")
    text = text.replace("（或 13.5）", "（亦即 13.5）")
    text = text.replace("(或 13.5)", "（亦即 13.5）")
    text = text.replace("（或13.5）", "（亦即 13.5）")
    text = text.replace("(或13.5)", "（亦即 13.5）")

    parts = re.split(r'(\([1-9]\))', text)
    for part in parts:
        if not part:
            continue
        if re.match(r'^\([1-9]\)$', part):
            r = paragraph.add_run(part + " ")
            set_run_font(r, font_name="Times New Roman", font_size=font_size, bold=True, color=default_color)
            continue

        tokens = re.split(r'([\u4e00-\u9fa5\u3000-\u303f\uff01-\uff60]+)', part)
        for tok in tokens:
            if not tok:
                continue

            if re.search(r'[\u4e00-\u9fa5\u3000-\u303f\uff01-\uff60]', tok):
                r = paragraph.add_run(tok)
                set_run_font(r, font_name="標楷體", font_size=font_size, bold=bold, color=default_color)
                continue

            tok_strip = tok.strip()
            if not tok_strip:
                paragraph.add_run(tok)
                continue

            if re.match(r'^\([A-D]\)$', tok_strip):
                r = paragraph.add_run(tok)
                set_run_font(r, font_name="Times New Roman", font_size=font_size, bold=True, color=default_color)
                continue

            has_math_op = any(c in tok_strip for c in ['=', '/', '^', '√', '*', '×', '°', ':', '△', '∠', '>', '<', '~', '∼', '∥'])
            has_arrow = "=>" in tok_strip or "⇒" in tok_strip
            is_var_expr = bool(re.search(r'[A-Za-z]', tok_strip) and re.search(r'[0-9\+\-\(\)]', tok_strip))
            is_geom_label = tok_strip in ["△ABC", "△DEF", "△ADE", "△COD", "BCED", "AB", "BC", "CA", "CD", "DE", "EF", "DF", "GH", "OA", "OB", "OC", "AG", "GD", "AD"]

            if has_math_op or has_arrow or is_var_expr or is_geom_label:
                l_sp = len(tok) - len(tok.lstrip())
                r_sp = len(tok) - len(tok.rstrip())
                if l_sp > 0:
                    paragraph.add_run(" " * l_sp)

                omml_content = parse_math_to_omml(tok_strip, font_size=font_size)
                if omml_content:
                    omath_xml = f'<m:oMath {MATH_NS}>{omml_content}</m:oMath>'
                    paragraph._p.append(parse_xml(omath_xml))
                else:
                    r = paragraph.add_run(tok_strip)
                    set_run_font(r, font_name="Times New Roman", font_size=font_size, bold=bold, color=default_color)

                if r_sp > 0:
                    paragraph.add_run(" " * r_sp)
            else:
                r = paragraph.add_run(tok)
                set_run_font(r, font_name="Times New Roman", font_size=font_size, bold=bold, color=default_color)

def append_math_to_cell(cell, text: str, font_size=9.5, bold=False, color=None, center=True):
    """在表格儲存格中寫入支援 OMML 方程式的排版內容"""
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    has_star = False
    clean_text = text.strip()
    if clean_text.startswith("* ") or clean_text.startswith("*"):
        has_star = True
        clean_text = clean_text.lstrip("*").strip()

    lines = [clean_text]
    if "(1)" in clean_text and "(2)" in clean_text:
        parts = clean_text.split("(2)")
        l1 = parts[0].strip()
        if l1.endswith(","):
            l1 = l1[:-1].strip()
        l2 = "(2) " + parts[1].strip()
        lines = [l1, l2]

    for idx, line in enumerate(lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            if center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if idx == 0 and has_star:
            r_star = p.add_run("* ")
            set_run_font(r_star, font_name="Times New Roman", font_size=font_size, bold=True, color=RGBColor(192, 0, 0))
        append_mixed_text(p, line, font_size=font_size, bold=bold, default_color=color)

def add_jinhe_header(doc, school_info="新北市立錦和高級中學", exam_title="115學年度第一學期國中部九年級數學科第三次段考試題", exam_range="九上 第 3 章 幾何與證明"):
    """錦和標準卷頭：主標題粗體底線、出題範圍、班級座號姓名、紅字扣5分警語與黑色墨水筆作答警語"""
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.left_indent = Pt(0)
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(3)
    title_p.paragraph_format.line_spacing = 1.15
    r_title = title_p.add_run(f"{school_info}{exam_title}")
    set_run_font(r_title, font_name="標楷體", font_size=13.5, bold=True, underline=True)

    info_p = doc.add_paragraph()
    info_p.paragraph_format.left_indent = Pt(0)
    info_p.paragraph_format.space_before = Pt(0)
    info_p.paragraph_format.space_after = Pt(3)
    info_p.paragraph_format.line_spacing = 1.15
    r_range = info_p.add_run(f"﹝命題範圍：{exam_range}﹞")
    set_run_font(r_range, font_name="標楷體", font_size=10.5, bold=True)
    info_p.add_run("   ")
    r_stu = info_p.add_run("班級：________  座號：____  姓名：____________")
    set_run_font(r_stu, font_name="標楷體", font_size=10.5, bold=True)

    w1_p = doc.add_paragraph()
    w1_p.paragraph_format.left_indent = Pt(0)
    w1_p.paragraph_format.space_before = Pt(0)
    w1_p.paragraph_format.space_after = Pt(2)
    w1_p.paragraph_format.line_spacing = 1.15
    r_w1 = w1_p.add_run("※ 答案卷(卡)未寫班級、姓名、座號，或畫卡錯誤致電腦無法判讀考生身份者，一律扣5分。")
    set_run_font(r_w1, font_name="標楷體", font_size=10, bold=True, color=RGBColor(200, 0, 0))

    w2_p = doc.add_paragraph()
    w2_p.paragraph_format.left_indent = Pt(0)
    w2_p.paragraph_format.space_before = Pt(0)
    w2_p.paragraph_format.space_after = Pt(4)
    w2_p.paragraph_format.line_spacing = 1.15
    r_w2 = w2_p.add_run("※ 本試卷所有試題（選擇題、填充題、非選擇題）請一律在答案卷上作答，請一律用黑色墨水筆作答，違者扣總分5分。")
    set_run_font(r_w2, font_name="標楷體", font_size=10, bold=True)

def create_question_p(doc, num_str, text_prefix, scope_str=None, space_after=2):
    """
    建立題號凸排 (Hanging Indent) 段落：
    - left_indent = Inches(0.28)
    - first_line_indent = -Inches(0.28)
    - 出題範圍 (3-X) 右對齊 7.48 英吋且取消粗體 (bold=False)
    """
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.48), WD_TAB_ALIGNMENT.RIGHT)

    r_num = p.add_run(num_str)
    set_run_font(r_num, font_name="Times New Roman", font_size=11, bold=True)
    r_txt = p.add_run(text_prefix)
    set_run_font(r_txt, font_name="標楷體", font_size=11)

    if scope_str:
        r_tab = p.add_run(f"\t{scope_str}")
        set_run_font(r_tab, font_name="Times New Roman", font_size=10.5, bold=False)

    return p

def create_options_p(doc, opt_text, space_after=3):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(opt_text)
    set_run_font(r, font_name="Times New Roman", font_size=11)
    return p

def safe_save(doc, filename):
    """Windows 檔案鎖定安全防禦存檔機制"""
    os.makedirs(os.path.dirname(os.path.abspath(filename)), exist_ok=True)
    try:
        doc.save(filename)
        return filename
    except PermissionError:
        base, ext = os.path.splitext(filename)
        alt_filename = f"{base}_修訂版{ext}"
        doc.save(alt_filename)
        print(f"[NOTE] 因 {filename} 正在被 Word 開啟鎖定，已儲存為：{alt_filename}")
        return alt_filename
